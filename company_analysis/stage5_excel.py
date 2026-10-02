# Stage 5: Save the results to a formatted Excel file.
#
# Uses the openpyxl library (install with: pip install openpyxl).

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from metrics import all_metrics, METRIC_KINDS
from stage6_credit import credit_rating

HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
BODY_FONT = Font(name="Arial")

# Excel number formats for each kind of value
EXCEL_FORMATS = {
    "pct": "0.0%",
    "x": '0.00"x"',
    "usd": "$#,##0",
    "usd2": "$#,##0.00",
}


def write_header(sheet, headings):
    for col, text in enumerate(headings, start=1):
        cell = sheet.cell(row=1, column=col, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)


def write_row(sheet, row, values, formats):
    for col, (value, kind) in enumerate(zip(values, formats), start=1):
        cell = sheet.cell(row=row, column=col, value=value)
        cell.font = BODY_FONT
        if kind:
            cell.number_format = EXCEL_FORMATS[kind]


def set_widths(sheet, widths):
    for col, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=col).column_letter].width = width


def save_report(companies, dcf_results, assumptions, filename):
    wb = Workbook()

    # Sheet 1: metrics for every company
    sheet = wb.active
    sheet.title = "Metrics"
    names = list(METRIC_KINDS.keys())
    write_header(sheet, ["Company", "Ticker", "Group", "Fiscal year end"] + names)
    for row, c in enumerate(companies, start=2):
        metrics = all_metrics(c)
        values = [c["name"], c["ticker"], c.get("group", ""), c.get("fiscal_year_end", "")]
        values += [metrics[n] for n in names]
        formats = [None, None, None, None] + [METRIC_KINDS[n] for n in names]
        write_row(sheet, row, values, formats)
    set_widths(sheet, [28, 8, 16, 14] + [14] * len(names))
    sheet.freeze_panes = "C2"

    # Sheet 2: DCF results
    sheet = wb.create_sheet("DCF")
    write_header(sheet, ["Company", "Ticker", "Starting FCF ($mm)", "Enterprise value ($mm)",
                         "Equity value ($mm)", "Terminal value % of EV", "Value per share"])
    for row, c in enumerate(companies, start=2):
        r = dcf_results[c["ticker"]]
        if r is None:
            write_row(sheet, row, [c["name"], c["ticker"], "Not enough positive FCF"], [None] * 3)
            continue
        write_row(sheet, row,
                  [c["name"], c["ticker"], r["starting_fcf"], r["enterprise_value"],
                   r["equity_value"], r["terminal_share"], r["value_per_share"]],
                  [None, None, "usd", "usd", "usd", "pct", "usd2"])
    last = len(companies) + 3
    sheet.cell(row=last, column=1, value="Assumptions").font = Font(name="Arial", bold=True)
    labels = [("FCF growth, years 1-5", "growth"), ("Discount rate", "discount"),
              ("Terminal growth", "terminal")]
    for i, (label, key) in enumerate(labels, start=1):
        write_row(sheet, last + i, [label, None, assumptions[key]], [None, None, "pct"])
    set_widths(sheet, [28, 8, 18, 20, 18, 18, 16])

    # Sheet 3: credit scorecard
    sheet = wb.create_sheet("Credit")
    write_header(sheet, ["Company", "Ticker", "Debt / EBITDA", "Interest coverage",
                         "Score (out of 8)", "Rating"])
    for row, c in enumerate(companies, start=2):
        metrics = all_metrics(c)
        rating, score = credit_rating(c)
        write_row(sheet, row,
                  [c["name"], c["ticker"], metrics["Debt / EBITDA"],
                   metrics["Interest coverage"], score, rating],
                  [None, None, "x", "x", None, None])
    set_widths(sheet, [28, 8, 14, 16, 14, 16])

    # Sheet 4: data notes, so readers know what was estimated
    sheet = wb.create_sheet("Notes")
    write_header(sheet, ["Ticker", "Note"])
    row = 2
    for c in companies:
        for note in c.get("notes", []) + [f"missing: {m}" for m in c.get("missing", [])]:
            write_row(sheet, row, [c["ticker"], note], [None, None])
            row += 1
    write_row(sheet, row + 1, ["Source", "SEC EDGAR XBRL company facts (latest 10-K). $ in millions."], [None, None])
    write_row(sheet, row + 2, ["Caution", "Asset managers consolidate funds/insurance units, so their "
                                          "leverage and margins are not comparable to operating companies."],
              [None, None])
    set_widths(sheet, [10, 90])

    wb.save(filename)
